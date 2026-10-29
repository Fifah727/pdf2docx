"""Tests de non-régression : un PDF synthétique par difficulté, vérifications structurelles sur le DOCX produit.

Lancer :  python -m pytest -q tests
(les PDF synthétiques sont régénérés par bench/make_corpus.py ; LaTeX requis pour les créer)
"""
import os
import sys

import pytest

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
sys.path.insert(0, ROOT)

from docx import Document  # noqa: E402
from docx.oxml.ns import qn  # noqa: E402

from pdf2docx import Config, convert  # noqa: E402

SYNTH = os.path.join(ROOT, "corpus", "synth")


def need(name):
    p = os.path.join(SYNTH, name + ".pdf")
    if not os.path.exists(p):
        pytest.skip("corpus absent : python bench/make_corpus.py")
    return p


def run(name, tmp_path, **kw):
    out = str(tmp_path / (name + ".docx"))
    pages = convert(need(name), out, Config(**kw) if kw else Config())
    return Document(out), pages


def text_of(doc):
    parts = [p.text for p in doc.paragraphs]
    for t in doc.tables:
        for row in t.rows:
            for c in row.cells:
                parts.append(c.text)
    return " ".join(parts)


def test_trois_colonnes_vraies_sections(tmp_path):
    doc, pages = run("s1_trois_colonnes", tmp_path)
    cols = [c.get(qn("w:num")) for c in doc.element.body.iter(qn("w:cols"))]
    assert "3" in cols, "une vraie section Word à 3 colonnes est attendue"
    assert "Introduction" in text_of(doc) and "Résultats" in text_of(doc)


def test_tableau_sans_bordure_detecte(tmp_path):
    doc, pages = run("s3_tableaux_sans_bordure", tmp_path)
    assert len(doc.tables) >= 2
    grid = [[c.text for c in r.cells] for r in doc.tables[0].rows]
    flat = " ".join(" ".join(r) for r in grid)
    assert "Nord" in flat and "1 250" in flat.replace(" ", " ").replace(" ", " ") or "1250" in flat


def test_cellules_fusionnees(tmp_path):
    doc, pages = run("s4_cellules_fusionnees", tmp_path)
    xml = doc.element.body.xml
    assert "w:gridSpan" in xml, "fusion horizontale attendue"
    assert "w:vMerge" in xml, "fusion verticale attendue"


def test_encadres_ont_un_fond(tmp_path):
    doc, pages = run("s2_encadres_image", tmp_path)
    assert pages[0].stats["boxes"] == 2
    assert "w:shd" in doc.element.body.xml


def test_formules_en_images(tmp_path):
    doc, pages = run("s5_formules", tmp_path)
    assert pages[0].stats["formulas"] >= 2
    assert len(doc.inline_shapes) >= 2


def test_scan_incline_ocr_et_tableau(tmp_path):
    doc, pages = run("s7_scan_fr_incline", tmp_path)
    st = pages[0].stats
    assert st["ocr_lines"] > 10
    assert abs(st["deskew_x100"] + 160) <= 30, "inclinaison de -1,6° attendue"
    assert st["tables_scanned"] == 1
    t = text_of(doc).lower()
    for mot in ("réunion", "reconstruire", "matériel", "services"):
        assert mot in t


def test_rapport_entetes_et_pied(tmp_path):
    doc, pages = run("s6_rapport_classique", tmp_path)
    sec = doc.sections[0]
    assert "Rapport 2026" in " ".join(p.text for p in sec.header.paragraphs) or \
        "Société Exemple" in " ".join(p.text for p in sec.header.paragraphs)
    assert "w:fldSimple" in sec.footer._element.xml, "numéro de page = vrai champ Word"


def test_nombre_de_pages_conserve(tmp_path):
    import pymupdf
    for name in ("s1_trois_colonnes", "s6_rapport_classique"):
        doc, pages = run(name, tmp_path)
        assert len(pages) == len(pymupdf.open(need(name)))


# --- régressions issues d'un document réel (examen scolaire : cadre, soulignements, forme courbe, drapeau transparent)
REAL = os.path.join(ROOT, "corpus", "real", "user_exam.pdf")


@pytest.mark.skipif(not os.path.exists(REAL), reason="corpus réel absent")
def test_document_reel_cadre_soulignements_image_tableau(tmp_path):
    import io
    from PIL import Image
    out = str(tmp_path / "exam.docx")
    pages = convert(REAL, out, Config())
    doc = Document(out)
    xml = doc.element.body.xml
    assert pages[0].stats["shapes"] >= 5, "cadre d'en-tête et forme courbe attendus en vectoriel"
    assert "wordprocessingShape" in xml
    assert xml.count("<w:u ") >= 4, "DREN / CISCO / ZAP / EPP / CLASSE DE soulignés (vrais soulignements)"
    # le drapeau garde sa transparence
    has_alpha = False
    for rel in doc.part.rels.values():
        if "image" in rel.reltype:
            im = Image.open(io.BytesIO(rel.target_part.blob))
            has_alpha |= im.mode in ("RGBA", "LA")
    assert has_alpha, "image à masque : canal alpha attendu"
    # pas de ligne fantôme : le tableau à bordures a 3 lignes (et non le titre « 1) Asio X »)
    t = [t for t in doc.tables if len(t.columns) == 3]
    assert t and len(t[0].rows) == 3


# --- régression : image rognée par un tracé de découpe + soulignement dessiné par un rectangle plein très fin
REAL2 = os.path.join(ROOT, "corpus", "real", "user_fifamangiana.pdf")


@pytest.mark.skipif(not os.path.exists(REAL2), reason="corpus réel absent")
def test_image_rognee_et_soulignement_rectangle(tmp_path):
    import io
    from PIL import Image
    out = str(tmp_path / "f.docx")
    convert(REAL2, out, Config())
    doc = Document(out)
    assert "<w:u " in doc.element.body.xml, "« LAHATSORATRA » doit rester souligné"
    for rel in doc.part.rels.values():
        if "image" in rel.reltype:
            w, h = Image.open(io.BytesIO(rel.target_part.blob)).size
            # partie visible : 498 x 365 pt (rapport 1,36) ; l'image complète fait 919 x 840 px (rapport 1,09)
            assert abs(w / h - 498.2 / 365.4) < 0.08, "l'image doit être recadrée sur sa partie visible"

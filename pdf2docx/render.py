"""Écriture du DOCX (python-docx + XML OOXML direct) : sections, colonnes, tableaux, encadrés, images, en-têtes."""
import io
from copy import deepcopy

from docx import Document
from docx.enum.table import WD_ROW_HEIGHT_RULE, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING, WD_BREAK, WD_TAB_ALIGNMENT
from docx.oxml import parse_xml, OxmlElement
from docx.oxml.ns import qn, nsdecls
from docx.shared import Pt, Emu, RGBColor
from PIL import Image as PILImage

from .fonts import style_of
from . import metrics
from .model import node_top

_TBLPR = ['tblStyle', 'tblpPr', 'tblOverlap', 'bidiVisual', 'tblStyleRowBandSize', 'tblStyleColBandSize',
          'tblW', 'jc', 'tblCellSpacing', 'tblInd', 'tblBorders', 'shd', 'tblLayout', 'tblCellMar',
          'tblLook', 'tblCaption', 'tblDescription']
A_L, A_C, A_R, A_J = (WD_ALIGN_PARAGRAPH.LEFT, WD_ALIGN_PARAGRAPH.CENTER,
                      WD_ALIGN_PARAGRAPH.RIGHT, WD_ALIGN_PARAGRAPH.JUSTIFY)


def tw(pt):
    return int(round(pt * 20))


def emu(pt):
    return int(round(pt * 12700))


class Ctr:
    """Conteneur d'écriture : corps du document ou cellule."""

    def __init__(self, obj, is_cell=False):
        self.obj = obj
        self.is_cell = is_cell
        self.reuse = obj.paragraphs[0] if is_cell else None
        self.last_tbl = False


class Sections:
    """Gère les ruptures de section en écrivant le sectPr de chaque section dans son dernier paragraphe."""

    def __init__(self, doc):
        self.doc = doc
        self.body = doc.element.body
        self.cur = None
        self.mark = 0
        self.refs = []
        self.first_closed = False

    def _count(self):
        return len(self.body) - 1       # le dernier enfant est le sectPr final

    def begin(self, **props):
        if self.cur is not None:
            if self._count() == self.mark:          # section vide : on fusionne
                props["start"] = self.cur["start"] if props.get("start") == "continuous" else props["start"]
                self.cur = props
                return
            self._close(self.cur, final=False)
        self.cur = props
        self.mark = self._count()

    def finish(self):
        if self._count() == self.mark:
            self.doc.add_paragraph()
            self._tiny_last()
        self._close(self.cur, final=True)

    def _tiny_last(self):
        p = self.doc.paragraphs[-1]
        pf = p.paragraph_format
        pf.space_before = pf.space_after = Pt(0)
        pf.line_spacing_rule = WD_LINE_SPACING.EXACTLY
        pf.line_spacing = Pt(1)

    def _build(self, p):
        W, H = p["W"], p["H"]
        cols = p.get("cols")
        if cols and len(cols) > 1:
            textw = W - p["left"] - p["right"]
            ws = [c[0] for c in cols]
            sps = [c[1] for c in cols]
            sps[-1] = 0
            ws[-1] = max(10, textw - sum(ws[:-1]) - sum(sps[:-1]))
            cx = f'<w:cols w:num="{len(cols)}" w:equalWidth="0">' + "".join(
                f'<w:col w:w="{tw(w)}" w:space="{tw(s)}"/>' for w, s in zip(ws, sps)) + '</w:cols>'
        else:
            cx = '<w:cols w:space="720"/>'
        xml = (f'<w:sectPr {nsdecls("w", "r")}><w:type w:val="{p["start"]}"/>'
               f'<w:pgSz w:w="{tw(W)}" w:h="{tw(H)}"/>'
               f'<w:pgMar w:top="{tw(p["top"])}" w:right="{tw(p["right"])}" w:bottom="{tw(p["bottom"])}" '
               f'w:left="{tw(p["left"])}" w:header="{tw(p.get("hdr", 20))}" w:footer="{tw(p.get("ftr", 20))}" w:gutter="0"/>'
               f'{cx}<w:docGrid w:linePitch="360"/></w:sectPr>')
        el = parse_xml(xml)
        if not self.first_closed:
            for r in reversed(self.refs):
                el.insert(0, deepcopy(r))
        return el

    def _close(self, props, final):
        el = self._build(props)
        self.first_closed = True
        if final:
            old = self.body.sectPr
            for r in reversed(self.refs):
                if not any(c.tag == r.tag and c.get(qn("r:id")) == r.get(qn("r:id")) for c in el):
                    el.insert(0, deepcopy(r))
            self.body.replace(old, el)
            return
        last = self.body[self._count() - 1]
        if last.tag != qn("w:p"):
            self.doc.add_paragraph()
            self._tiny_last()
            last = self.body[self._count() - 1]
        last.set_sectPr(el)


class Renderer:
    def __init__(self, cfg):
        self.cfg = cfg
        self.doc = Document()
        self.secs = Sections(self.doc)
        self.pending_break = False
        self.pm = None
        self.forced_breaks = 0
        self._shape_id = 0
        self.in_hf = False
        self.ftr_top = None
        self.hdr_dist = self.ftr_dist = 20.0
        self.page_top, self.page_bottom = 36.0, 14.0
        self.doc.core_properties.author = "pdf2docx"
        st = self.doc.styles["Normal"]
        st.font.name = "Arial"
        st.font.size = Pt(10)

    # ------------------------------------------------------------------ conteneurs
    def par(self, ctr):
        if ctr.reuse is not None:
            p = ctr.reuse
            ctr.reuse = None
        else:
            p = ctr.obj.add_paragraph()
        ctr.last_tbl = False
        if self.pending_break:
            # le saut de colonne vit dans son propre paragraphe minuscule (sinon space_before compté 2 fois)
            self.pending_break = False
            bp = p
            bp.add_run().add_break(WD_BREAK.COLUMN)
            self._tiny_para(bp)
            p = ctr.obj.add_paragraph()
        return p

    def _tiny_mark(self, p, sz=2):
        pPr = p._p.get_or_add_pPr()
        for old in pPr.findall(qn("w:rPr")):
            pPr.remove(old)
        rPr = OxmlElement("w:rPr")
        s = OxmlElement("w:sz")
        s.set(qn("w:val"), str(sz))
        rPr.append(s)
        pPr.append(rPr)

    def spacer(self, ctr, h):
        p = self.par(ctr)
        pf = p.paragraph_format
        pf.space_before = pf.space_after = Pt(0)
        pf.line_spacing_rule = WD_LINE_SPACING.EXACTLY
        pf.line_spacing = Pt(max(1.0, min(h, 1500)))
        pf.widow_control = False
        self._tiny_mark(p)
        return p

    def _new_table(self, ctr, rows, cols, width_pt):
        if ctr.reuse is not None:                      # supprime le paragraphe vide initial de la cellule
            el = ctr.reuse._p
            el.getparent().remove(el)
            ctr.reuse = None
        if self.pending_break:
            self.par(ctr)
            self.spacer(ctr, 1)
        if ctr.last_tbl:
            self.spacer(ctr, 1)
        if ctr.is_cell:
            tbl = ctr.obj.add_table(rows, cols)
            ctr.reuse = ctr.obj.paragraphs[-1]          # paragraphe final obligatoire, réutilisable
            self._tiny_para(ctr.reuse)
        else:
            tbl = ctr.obj.add_table(rows, cols)
        ctr.last_tbl = True
        return tbl

    def _tiny_para(self, p):
        pf = p.paragraph_format
        pf.space_before = pf.space_after = Pt(0)
        pf.line_spacing_rule = WD_LINE_SPACING.EXACTLY
        pf.line_spacing = Pt(1)
        self._tiny_mark(p)

    @staticmethod
    def _tblpr_put(tblPr, el):
        tag = el.tag.split("}")[1]
        for old in tblPr.findall(qn("w:" + tag)):
            tblPr.remove(old)
        succ = {qn("w:" + t) for t in _TBLPR[_TBLPR.index(tag) + 1:]}
        for i, ch in enumerate(list(tblPr)):
            if ch.tag in succ:
                tblPr.insert(i, el)
                return
        tblPr.append(el)

    def _setup_table(self, table, width, indent, borders=None):
        tblPr = table._tbl.tblPr
        self._tblpr_put(tblPr, parse_xml(f'<w:tblW {nsdecls("w")} w:w="{tw(width)}" w:type="dxa"/>'))
        self._tblpr_put(tblPr, parse_xml(f'<w:tblInd {nsdecls("w")} w:w="{tw(indent)}" w:type="dxa"/>'))
        if borders:
            col, wpt = borders
            sz = max(2, int(round(wpt * 8)))
            b = "".join(f'<w:{s} w:val="single" w:sz="{sz}" w:space="0" w:color="{col}"/>'
                        for s in ("top", "left", "bottom", "right", "insideH", "insideV"))
            self._tblpr_put(tblPr, parse_xml(f'<w:tblBorders {nsdecls("w")}>{b}</w:tblBorders>'))
        self._tblpr_put(tblPr, parse_xml(f'<w:tblLayout {nsdecls("w")} w:type="fixed"/>'))
        self._tblpr_put(tblPr, parse_xml(
            f'<w:tblCellMar {nsdecls("w")}><w:top w:w="0" w:type="dxa"/><w:left w:w="0" w:type="dxa"/>'
            f'<w:bottom w:w="0" w:type="dxa"/><w:right w:w="0" w:type="dxa"/></w:tblCellMar>'))

    @staticmethod
    def _cell_borders(cell, borders):
        if not borders:
            return
        tcPr = cell._tc.get_or_add_tcPr()
        for old in tcPr.findall(qn("w:tcBorders")):
            tcPr.remove(old)
        xml = "".join(
            f'<w:{side} w:val="single" w:sz="{max(2, int(round(borders[side][1] * 8)))}" w:space="0" w:color="{borders[side][0]}"/>'
            for side in ("top", "left", "bottom", "right") if side in borders)
        el = parse_xml(f'<w:tcBorders {nsdecls("w")}>{xml}</w:tcBorders>')
        tcPr.insert_element_before(el, "w:shd", "w:noWrap", "w:tcMar", "w:textDirection", "w:tcFitText", "w:vAlign", "w:hideMark")

    @staticmethod
    def _shade(cell, hexfill):
        tcPr = cell._tc.get_or_add_tcPr()
        for old in tcPr.findall(qn("w:shd")):
            tcPr.remove(old)
        shd = parse_xml(f'<w:shd {nsdecls("w")} w:val="clear" w:color="auto" w:fill="{hexfill}"/>')
        tcPr.insert_element_before(shd, "w:noWrap", "w:tcMar", "w:textDirection", "w:tcFitText",
                                   "w:vAlign", "w:hideMark")

    # ------------------------------------------------------------------ runs
    def _style_run(self, run, fam, bold, ital, size, color, script, spacing=0.0, under=False):
        f = run.font
        if under:
            f.underline = True
        f.name = fam
        rpr = run._r.get_or_add_rPr()
        rf = rpr.find(qn("w:rFonts"))
        if rf is not None:
            rf.set(qn("w:eastAsia"), fam)
            rf.set(qn("w:cs"), fam)
        f.size = Pt(size)
        f.bold = bool(bold)
        f.italic = bool(ital)
        f.color.rgb = RGBColor((color >> 16) & 255, (color >> 8) & 255, color & 255)
        if script == "sup":
            f.superscript = True
        elif script == "sub":
            f.subscript = True
        if abs(spacing) >= 0.05:
            sp = OxmlElement("w:spacing")
            sp.set(qn("w:val"), str(int(round(spacing * 20))))
            rpr.insert_element_before(sp, "w:w", "w:kern", "w:position", "w:sz", "w:szCs", "w:highlight",
                                      "w:u", "w:effect", "w:bdr", "w:shd", "w:fitText", "w:vertAlign",
                                      "w:rtl", "w:cs", "w:em", "w:lang", "w:eastAsianLayout", "w:specVanish", "w:oMath")

    def _fit(self, fam, b, i, size, w):
        """(largeur attendue, largeur mesurée, nb de caractères) d'un mot avec la police de remplacement."""
        pred = metrics.width(fam, b, i, size, w.text)
        if pred is None or len(w.text) < 1:
            return 0.0, 0.0, 0
        return (w.x1 - w.x0), pred, len(w.text)

    def _tokens(self, p, keep_breaks, skip_marker=True, avail=None):
        """Découpe le paragraphe en séries homogènes. `avail` = (largeur 1re ligne, largeur autres lignes) :
        si fourni, on simule la césure de Word et on force un saut de ligne là où Word remonterait un mot."""
        toks = []
        first = True
        prev_word = None
        line_w = 0.0
        forced = 0
        for li, ln in enumerate(p.lines):
            for wi, w in enumerate(ln.words):
                if skip_marker and p.marker is w:
                    continue
                fam, b, i = style_of(w.font, w.flags)
                size = round((ln.size if w.script else w.size) * 2) / 2
                wexp = w.x1 - w.x0
                spw = (metrics.width(fam, b, i, size, " ") or 0.25 * size)
                sep = ""
                if not first:
                    if wi == 0:
                        hyph = (prev_word is not None and prev_word.text.endswith("-") and len(prev_word.text) > 2
                                and prev_word.text[-2].isalpha() and w.text[:1].islower())
                        if keep_breaks:
                            sep = "\n"
                        elif hyph:
                            toks[-1][0] = toks[-1][0][:-1] + "\u00ad"   # trait d'union conditionnel : même coupure que le PDF
                            sep = ""
                            line_w += wexp          # le mot continue sur la ligne suivante
                        else:
                            sep = " "
                            if avail is not None:
                                room = avail[0] if li - 1 == 0 else avail[1]
                                if line_w + spw + wexp <= room:         # Word remonterait ce mot : on force la coupure
                                    sep = "\n"
                                    forced += 1
                        if not (not keep_breaks and hyph):
                            line_w = wexp
                    else:
                        nsp = self._gap_spaces(ln.words[wi - 1], w, spw) if w.space else 0
                        sep = " " * nsp
                        line_w += nsp * spw + wexp
                else:
                    line_w = wexp
                key = (fam, b, i, size, w.color, w.script, w.under)
                exp, pred, nch = self._fit(fam, b, i, size if not w.script else w.size, w)
                if toks and toks[-1][1] == key:
                    toks[-1][0] += sep + w.text
                    acc = toks[-1][2]
                else:
                    acc = [0.0, 0.0, 0]
                    if sep and key[6] and toks and not toks[-1][1][6]:
                        toks[-1][0] += sep          # l'espace / saut avant un mot souligné reste hors du soulignement
                        sep = ""
                    toks.append([sep + w.text, key, acc])
                acc[0] += exp
                acc[1] += pred
                acc[2] += nch
                first = False
                prev_word = w
        self.forced_breaks += forced
        return toks

    @staticmethod
    def _gap_spaces(prev, w, spw):
        """Nombre d'espaces à écrire entre deux mots : 1, ou plusieurs si le PDF laisse un grand blanc (tabulation, colonnes)."""
        gap = w.x0 - prev.x1
        if gap > 4.5 * spw and gap > 14.0 and spw > 0.5:
            return int(min(80, max(1, round(gap / spw))))
        return 1

    def _need(self, p, avail_other, first):
        """Excès maximal (pt) de la largeur naturelle d'une ligne sur la largeur disponible."""
        worst = -1e9
        for li, ln in enumerate(p.lines):
            wsum = 0.0
            for wi, w in enumerate(ln.words):
                if p.marker is w:
                    continue
                fam, b, i = style_of(w.font, w.flags)
                size = round((ln.size if w.script else w.size) * 2) / 2
                wsum += (w.x1 - w.x0)
                if wi > 0 and w.space:
                    spw = metrics.width(fam, b, i, size, " ") or 0.25 * size
                    wsum += spw * self._gap_spaces(ln.words[wi - 1], w, spw)
            room = avail_other - (first if li == 0 else 0.0)
            worst = max(worst, wsum - room)
        return worst

    @staticmethod
    def _spacing(acc):
        exp, pred, n = acc
        if n < 2 or exp <= 0:
            return 0.0
        sp = (exp - pred) / n
        return max(-1.2, min(3.0, sp)) if abs(sp) >= 0.04 else 0.0

    def _emit_runs(self, par, toks, shrink=0.0):
        for text, (fam, b, i, size, color, script, under), acc in toks:
            sp = self._spacing(acc) - shrink
            parts = text.split("\n")
            for k, part in enumerate(parts):
                if k > 0:
                    par.add_run().add_break()
                if part:
                    r = par.add_run(part)
                    self._style_run(r, fam, b, i, size, color, script, sp, under)

    # ------------------------------------------------------------------ paragraphes
    def _place(self, p, L, R):
        if p.eff_L is not None:
            L = max(L, p.eff_L)
        if p.eff_R is not None:
            R = min(R, p.eff_R)
        W = max(R - L, 10)
        x0, x1, n = p.x0, p.x1, p.n
        lg, rg = x0 - L, R - x1
        align = A_L
        if n >= 3 and p.marker is None and all(l.x1 >= x1 - 1.5 for l in p.lines[:-1]) \
                and lg < 0.5 * W and (x1 - x0) > 0.5 * W:
            align = A_J
        elif all(abs((l.x0 + l.x1) / 2 - (x0 + x1) / 2) <= 2.0 for l in p.lines) and \
                abs((x0 + x1) / 2 - (L + R) / 2) <= 0.03 * W + 2 and lg > 0.06 * W:
            align = A_C
        elif all(abs(l.x1 - x1) <= 1.5 for l in p.lines) and rg < 0.03 * W + 3 and lg > 0.15 * W:
            align = A_R
        left = right = first = 0.0
        if align == A_C:
            pass
        elif align == A_R:
            right = max(rg, 0)
        else:
            if p.marker is not None and p.text_x0 is not None:
                left = p.text_x0 - L
                first = p.marker.x0 - p.text_x0
            elif n >= 2 and not p.near_float:
                others = min(l.x0 for l in p.lines[1:])
                left = others - L
                first = p.lines[0].x0 - others
                if abs(first) < 2:
                    first = 0.0
            else:
                left = lg
            right = rg if n >= 2 else 0.0     # tolérance ajoutée ensuite, calculée d'après les largeurs réelles
        return align, max(left, -60), right, first

    def emit_para(self, ctr, p, L, R, cursor):
        gap = p.top - cursor
        sb = round(max(0.0, gap) * 20) / 20
        par = self.par(ctr)
        pf = par.paragraph_format
        if p.level:
            par.style = self.doc.styles[f"Heading {min(p.level, 3)}"]
        align, left, right, first = self._place(p, L, R)
        par.alignment = align
        shrink = 0.0
        if p.n == 1 and ctr.is_cell:
            Wc = max((R if p.eff_R is None else min(R, p.eff_R)) - L, 10)
            over = self._need(p, Wc - left - max(right, 0.0), first)
            if over > -1.0:                    # le mot dépasse sa cellule : resserrement imperceptible pour qu'il tienne
                nch = max(sum(len(w.text) for w in p.words), 1)
                shrink = min(1.0, (over + 1.0) / nch)
        if p.n == 1 and (ctr.is_cell or self.in_hf):      # une seule ligne dans une cellule : jamais de repli (Word mesure parfois +1 pt)
            if align == A_C:
                left, right = left - 30.0, right - 30.0
            elif align == A_R:
                left = left - 40.0
            else:
                right = right - 40.0
        pf.left_indent = Pt(left)
        pf.right_indent = Pt(right)
        pf.first_line_indent = Pt(first)
        pf.space_before = Pt(sb)
        pf.space_after = Pt(0)
        pf.widow_control = False
        pf.keep_with_next = False
        pf.keep_together = False
        if self.cfg.line_mode == "exact":
            pf.line_spacing_rule = WD_LINE_SPACING.EXACTLY
            pf.line_spacing = Pt(p.pitch)
        else:
            pf.line_spacing = 1.0
        if p.bg:
            par._p.get_or_add_pPr().append(parse_xml(f'<w:shd {nsdecls("w")} w:val="clear" w:color="auto" w:fill="{p.bg}"/>'))
        if not ctr.is_cell:
            for f in p.floats:
                self._anchor_picture(par, f)
        if p.marker is not None:
            fam, b, i = style_of(p.marker.font, p.marker.flags)
            r = par.add_run(p.marker.text + "\t")
            self._style_run(r, fam, b, i, round(p.marker.size * 2) / 2, p.marker.color, None)
        W = max((R if p.eff_R is None else min(R, p.eff_R)) - (L if p.eff_L is None else max(L, p.eff_L)), 10)
        avail = None
        if p.n > 1 and align not in (A_C, A_R):
            need = self._need(p, W - left - right, first)          # dépassement max d'une ligne (police de remplacement)
            extra = max(self.cfg.slack * W * 0.25, need + 0.4)
            right -= extra
            pf.right_indent = Pt(right)
            if self.cfg.line_breaks == "sim":
                avail = (W - left - first - right, W - left - right)
        self._emit_runs(par, self._tokens(p, self.cfg.line_breaks == "keep", avail=avail), shrink)
        return cursor + sb + p.n * p.pitch

    # ------------------------------------------------------------------ tracés vectoriels
    def _anchor_shape(self, par, sh, idx):
        pad = 1.0
        x0, y0 = sh.x0 - pad, sh.y0 - pad
        w, h = max(sh.w, 0.0) + 2 * pad, max(sh.h, 0.0) + 2 * pad
        pw, ph = emu(w), emu(h)

        def pt(x, y):
            return f'<a:pt x="{emu(x - x0)}" y="{emu(y - y0)}"/>'
        cmds = []
        for c in sh.paths:
            if c[0] == "M":
                cmds.append(f'<a:moveTo>{pt(c[1], c[2])}</a:moveTo>')
            elif c[0] == "L":
                cmds.append(f'<a:lnTo>{pt(c[1], c[2])}</a:lnTo>')
            elif c[0] == "C":
                cmds.append(f'<a:cubicBezTo>{pt(c[1], c[2])}{pt(c[3], c[4])}{pt(c[5], c[6])}</a:cubicBezTo>')
            else:
                cmds.append('<a:close/>')
        fill = (f'<a:solidFill><a:srgbClr val="{sh.fill}"/></a:solidFill>' if sh.fill else '<a:noFill/>')
        ln = (f'<a:ln w="{emu(sh.width)}" cap="flat"><a:solidFill><a:srgbClr val="{sh.stroke}"/></a:solidFill>'
              + ('<a:prstDash val="dash"/>' if sh.dash else '') + '<a:miter lim="800000"/></a:ln>') if sh.stroke else '<a:ln><a:noFill/></a:ln>'
        xml = (
            f'<wp:anchor {nsdecls("wp")} distT="0" distB="0" distL="0" distR="0" simplePos="0" relativeHeight="{1000 + idx}" '
            f'behindDoc="1" locked="0" layoutInCell="1" allowOverlap="1">'
            f'<wp:simplePos x="0" y="0"/>'
            f'<wp:positionH relativeFrom="page"><wp:posOffset>{emu(x0)}</wp:posOffset></wp:positionH>'
            f'<wp:positionV relativeFrom="page"><wp:posOffset>{emu(y0)}</wp:posOffset></wp:positionV>'
            f'<wp:extent cx="{pw}" cy="{ph}"/><wp:effectExtent l="0" t="0" r="0" b="0"/><wp:wrapNone/>'
            f'<wp:docPr id="{5000 + idx}" name="Forme {idx}"/><wp:cNvGraphicFramePr/>'
            f'<a:graphic xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main">'
            f'<a:graphicData uri="http://schemas.microsoft.com/office/word/2010/wordprocessingShape">'
            f'<wps:wsp xmlns:wps="http://schemas.microsoft.com/office/word/2010/wordprocessingShape">'
            f'<wps:cNvSpPr/><wps:spPr><a:xfrm><a:off x="0" y="0"/><a:ext cx="{pw}" cy="{ph}"/></a:xfrm>'
            f'<a:custGeom><a:avLst/><a:gdLst/><a:ahLst/><a:cxnLst/><a:rect l="0" t="0" r="r" b="b"/>'
            f'<a:pathLst><a:path w="{pw}" h="{ph}"{"" if sh.fill else " fill=" + chr(34) + "none" + chr(34)}>{"".join(cmds)}</a:path></a:pathLst></a:custGeom>'
            f'{fill}{ln}</wps:spPr><wps:bodyPr/></wps:wsp></a:graphicData></a:graphic></wp:anchor>')
        run = par.add_run()
        drawing = OxmlElement("w:drawing")
        drawing.append(parse_xml(xml))
        run._r.append(drawing)

    # ------------------------------------------------------------------ images
    @staticmethod
    def _img_bytes(img):
        ext = (img.ext or "png").lower()
        if ext in ("png", "jpeg", "jpg", "gif", "bmp", "tiff"):
            return img.data
        try:
            im = PILImage.open(io.BytesIO(img.data))
            if im.mode not in ("RGB", "RGBA", "L"):
                im = im.convert("RGB")
            buf = io.BytesIO()
            im.save(buf, "PNG")
            return buf.getvalue()
        except Exception:
            return None

    def _anchor_picture(self, par, img):
        data = self._img_bytes(img)
        if not data:
            return
        run = par.add_run()
        shape = run.add_picture(io.BytesIO(data), width=Pt(img.w), height=Pt(img.h))
        inline = shape._inline
        extent, docPr = inline.find(qn("wp:extent")), inline.find(qn("wp:docPr"))
        cnv, graphic = inline.find(qn("wp:cNvGraphicFramePr")), inline.find(qn("a:graphic"))
        anchor = parse_xml(
            f'<wp:anchor {nsdecls("wp")} distT="76200" distB="76200" distL="101600" distR="101600" '
            f'simplePos="0" relativeHeight="251658240" behindDoc="1" locked="0" layoutInCell="1" allowOverlap="1">'
            f'<wp:simplePos x="0" y="0"/>'
            f'<wp:positionH relativeFrom="page"><wp:posOffset>{emu(img.x0)}</wp:posOffset></wp:positionH>'
            f'<wp:positionV relativeFrom="page"><wp:posOffset>{emu(img.y0)}</wp:posOffset></wp:positionV></wp:anchor>')
        anchor.append(deepcopy(extent))
        anchor.append(parse_xml(f'<wp:effectExtent {nsdecls("wp")} l="0" t="0" r="0" b="0"/>'))
        anchor.append(parse_xml(f'<wp:wrapNone {nsdecls("wp")}/>'))
        anchor.append(deepcopy(docPr))
        if cnv is not None:
            anchor.append(deepcopy(cnv))
        anchor.append(deepcopy(graphic))
        inline.getparent().replace(inline, anchor)

    def emit_img(self, ctr, im, L, R, cursor):
        data = self._img_bytes(im)
        if not data:
            return cursor
        gap = im.top - cursor
        if gap > 0.5:
            self.spacer(ctr, gap)
        par = self.par(ctr)
        pf = par.paragraph_format
        pf.space_before = pf.space_after = Pt(0)
        pf.line_spacing = 1.0
        pf.widow_control = False
        w, h = im.w, im.h
        avail = max(R - L, 10)
        if w > avail + 1:
            h *= avail / w
            w = avail
        room = self.pm.H - self.page_bottom - max(cursor, im.top) - 4.0
        if h > room > 40:                      # jamais plus haut que la place restante (évite une page parasite)
            w *= room / h
            h = room
        lg, rg = im.x0 - L, R - im.x1
        if abs(lg - rg) <= 0.04 * avail + 2 and min(lg, rg) > 0.06 * avail:
            par.alignment = A_C
        elif rg < lg - 0.1 * avail:
            par.alignment = A_R
            pf.right_indent = Pt(max(rg, 0))
        else:
            par.alignment = A_L
            pf.left_indent = Pt(max(lg, 0))
        run = par.add_run()
        run.font.size = Pt(1)
        shape = run.add_picture(io.BytesIO(data), width=Pt(w), height=Pt(h))
        for k in ("distT", "distB", "distL", "distR"):       # sans ces attributs, LibreOffice ajoute 9 pt de marge autour de l'image
            shape._inline.set(k, "0")
        self._tiny_mark(par)
        return max(cursor, im.top) + h

    # ------------------------------------------------------------------ tableaux
    def emit_tbl(self, ctr, t, L, R, cursor):
        gap = t.top - cursor
        if gap > 0.5:
            self.spacer(ctr, gap)
        nr, nc = len(t.row_edges) - 1, len(t.col_edges) - 1
        width = t.col_edges[-1] - t.col_edges[0]
        table = self._new_table(ctr, nr, nc, width)
        self._setup_table(table, width, t.col_edges[0] - L, None)
        table.autofit = False
        for c in range(nc):
            table.columns[c].width = Emu(emu(t.col_edges[c + 1] - t.col_edges[c]))
        for r in range(nr):
            row = table.rows[r]
            row.height = Pt(max(1.0, t.row_edges[r + 1] - t.row_edges[r]))
            # tableau à traits : géométrie connue => hauteur exacte (sinon les bordures allongent chaque ligne et la page déborde)
            row.height_rule = WD_ROW_HEIGHT_RULE.EXACTLY if t.ruled else WD_ROW_HEIGHT_RULE.AT_LEAST
        # fusions d'abord, contenu ensuite
        handles = []
        for c in t.cells:
            tc = table.cell(c.r0, c.c0)
            if c.r1 - c.r0 > 1 or c.c1 - c.c0 > 1:
                tc = tc.merge(table.cell(c.r1 - 1, c.c1 - 1))
            tc.width = Emu(emu(t.col_edges[c.c1] - t.col_edges[c.c0]))
            handles.append(tc)
        for c, tc in zip(t.cells, handles):
            self._cell_borders(tc, c.borders)
            if c.fill:
                self._shade(tc, c.fill)
            if c.vcenter:
                tc.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            sub = Ctr(tc, is_cell=True)
            x0, y0, x1, y1 = c.rect
            start = c.paras[0].top if (c.vcenter and c.paras) else y0
            self.render_nodes(sub, c.paras, x0, x1, start)
            self._finish_cell(sub)
        return max(cursor, t.top) + (t.bottom - t.top)

    def _finish_cell(self, sub):
        if sub.reuse is not None:
            self._tiny_para(sub.reuse)

    # ------------------------------------------------------------------ encadrés
    def emit_box(self, ctr, b, L, R, cursor):
        gap = b.top - cursor
        if gap > 0.5:
            self.spacer(ctr, gap)
        table = self._new_table(ctr, 1, 1, b.w)
        self._setup_table(table, b.w, b.x0 - L, None)
        table.autofit = False
        table.columns[0].width = Emu(emu(b.w))
        row = table.rows[0]
        row.height = Pt(max(1.0, b.h))
        row.height_rule = WD_ROW_HEIGHT_RULE.AT_LEAST
        cell = table.cell(0, 0)
        cell.width = Emu(emu(b.w))
        if b.stroke:
            self._cell_borders(cell, {k: (b.stroke, b.stroke_w) for k in ("top", "bottom", "left", "right")})
        if b.fill:
            self._shade(cell, b.fill)
        sub = Ctr(cell, is_cell=True)
        self.render_nodes(sub, b.children, b.x0, b.x1, b.y0)
        self._finish_cell(sub)
        return max(cursor, b.top) + b.h

    # ------------------------------------------------------------------ colonnes
    def emit_cols_table(self, ctr, n, L, R, cursor):
        gap = n.top - cursor
        if gap > 0.5:
            self.spacer(ctr, gap)
        gb = n.gboxes
        edges = [gb[0][0]] + [(a[2] + b[0]) / 2 for a, b in zip(gb, gb[1:])] + [gb[-1][2]]
        width = edges[-1] - edges[0]
        table = self._new_table(ctr, 1, len(gb), width)
        self._setup_table(table, width, edges[0] - L, None)
        table.autofit = False
        for i in range(len(gb)):
            table.columns[i].width = Emu(emu(edges[i + 1] - edges[i]))
            table.cell(0, i).width = Emu(emu(edges[i + 1] - edges[i]))
        ends = []
        for i, g in enumerate(n.groups):
            sub = Ctr(table.cell(0, i), is_cell=True)
            ends.append(self.render_nodes(sub, g, edges[i], edges[i + 1], n.top))
            self._finish_cell(sub)
        return max(ends + [cursor])

    def emit_cols_real(self, ctr, n, L, R, cursor):
        pm = self.pm
        gb = n.gboxes
        cols = []
        for i, b in enumerate(gb):
            sp = gb[i + 1][0] - b[2] if i + 1 < len(gb) else 0
            cols.append((b[2] - b[0], max(sp, 0)))
        base = dict(W=pm.W, H=pm.H, top=self.page_top, bottom=self.page_bottom, hdr=self.hdr_dist, ftr=self.ftr_dist)
        self.secs.begin(start="continuous", left=gb[0][0], right=pm.W - gb[-1][2], cols=cols, **base)
        ends = []
        for i, g in enumerate(n.groups):
            if i > 0:
                self.pending_break = True
            gtop = min(node_top(x) for x in g) if g else cursor
            lead = gtop - cursor
            if lead > 0.5:
                self.spacer(ctr, lead)      # espaceur à hauteur fixe : indépendant du traitement de space_before en tête de colonne
            elif self.pending_break:
                self.spacer(ctr, 1)
            ends.append(self.render_nodes(ctr, g, gb[i][0], gb[i][2], max(cursor, gtop)))
        self.pending_break = False
        self.secs.begin(start="continuous", left=L, right=pm.W - R, cols=None, **base)
        return max(ends + [cursor])

    # ------------------------------------------------------------------ boucle
    @staticmethod
    def _floats_in(n):
        k = n.kind
        if k == "para":
            return list(n.floats)
        out = []
        if k == "cols":
            for g in n.groups:
                for x in g:
                    out += Renderer._floats_in(x)
        elif k == "tbl":
            for c in n.cells:
                for x in c.paras:
                    out += Renderer._floats_in(x)
        elif k == "box":
            for x in n.children:
                out += Renderer._floats_in(x)
        return out

    def render_nodes(self, ctr, nodes, L, R, cursor):
        for n in nodes:
            k = n.kind
            if k != "para" and not ctr.is_cell:
                fl = self._floats_in(n)
                if fl:                       # une image ancrée dans une cellule serait positionnée par rapport à la cellule
                    hp = self.par(ctr)
                    self._tiny_para(hp)
                    for f in fl:
                        self._anchor_picture(hp, f)
                    cursor += 1.0            # hauteur du paragraphe d'ancrage, déduite de l'espacement qui suit
            if k == "para":
                cursor = self.emit_para(ctr, n, L, R, cursor)
            elif k == "img":
                cursor = self.emit_img(ctr, n, L, R, cursor)
            elif k == "tbl":
                cursor = self.emit_tbl(ctr, n, L, R, cursor)
            elif k == "box":
                cursor = self.emit_box(ctr, n, L, R, cursor)
            elif k == "cols":
                if n.real and not ctr.is_cell:
                    cursor = self.emit_cols_real(ctr, n, L, R, cursor)
                else:
                    cursor = self.emit_cols_table(ctr, n, L, R, cursor)
        return cursor

    # ------------------------------------------------------------------ en-têtes / pieds
    def make_hf(self, hdr, ftr, npages, W, H, L, R):
        """hdr / ftr : listes de (n° de page d'origine, Para)."""
        sec = self.doc.sections[0]
        self.hdr_dist, self.ftr_dist = 20.0, 20.0
        for zone, items in (("h", hdr), ("f", ftr)):
            if not items:
                continue
            part = sec.header if zone == "h" else sec.footer
            part.is_linked_to_previous = False
            imgs = [it for it in items if it[1].kind == "img"]
            items = [it for it in items if it[1].kind == "para"]
            if imgs and not items:
                part.paragraphs[0].paragraph_format.space_after = Pt(0)
            for _, p in items:
                p.base_y = p.lines[0].base
            rows = []
            for pno, p in sorted(items, key=lambda z: z[1].base_y):
                if rows and abs(rows[-1][0][1].base_y - p.base_y) < 3:
                    rows[-1].append((pno, p))
                else:
                    rows.append([(pno, p)])
            first_par = True
            for row in rows:
                row.sort(key=lambda z: z[1].x0)
                par = part.paragraphs[0] if first_par else part.add_paragraph()
                first_par = False
                pf = par.paragraph_format
                pf.space_before = pf.space_after = Pt(0)
                pf.line_spacing_rule = WD_LINE_SPACING.EXACTLY
                pf.line_spacing = Pt(row[0][1].pitch)
                pf.left_indent = Pt(max(row[0][1].x0 - L, 0))
                pf.right_indent = Pt(-80)
                for k, (pno, p) in enumerate(row):
                    if k > 0:
                        right = abs(p.x1 - R) < 6
                        pos = (p.x1 if right else p.x0) - L
                        pf.tab_stops.add_tab_stop(Pt(pos), WD_TAB_ALIGNMENT.RIGHT if right else WD_TAB_ALIGNMENT.LEFT)
                        par.add_run("\t")
                    self._hf_runs(par, p, pno, npages)
            for _, im in imgs:                       # images répétées : ancrées à leur position absolue dans l'en-tête / le pied
                self._anchor_picture(part.paragraphs[0], im)
            ys = [p for _, p in items + imgs]
            if not ys:
                continue
            if zone == "h":
                self.hdr_dist = max(2.0, min(p.y0 for p in ys))
            else:
                self.ftr_dist = max(2.0, H - max(p.y1 for p in ys))
                self.ftr_top = min(p.y0 for p in ys)
        self.secs.refs = [deepcopy(e) for e in self.doc.element.body.sectPr
                          if e.tag in (qn("w:headerReference"), qn("w:footerReference"))]

    def _hf_runs(self, par, p, pno, npages):
        first = True
        for ln in p.lines:
            for w in ln.words:
                fam, b, i = style_of(w.font, w.flags)
                size = round(w.size * 2) / 2
                sep = "" if first else " "
                first = False
                if sep:
                    r = par.add_run(sep)
                    self._style_run(r, fam, b, i, size, w.color, None)
                if w.text.isdigit() and int(w.text) in (pno, npages):
                    instr = "PAGE" if int(w.text) == pno else "NUMPAGES"
                    self._add_field(par, instr, w.text, fam, b, i, size, w.color)
                else:
                    r = par.add_run(w.text)
                    self._style_run(r, fam, b, i, size, w.color, None, self._spacing(self._fit(fam, b, i, size, w)))

    def _add_field(self, par, instr, shown, fam, b, i, size, color):
        rpr = (f'<w:rPr><w:rFonts w:ascii="{fam}" w:hAnsi="{fam}" w:cs="{fam}"/>'
               f'{"<w:b/>" if b else ""}{"<w:i/>" if i else ""}'
               f'<w:color w:val="{color:06X}"/><w:sz w:val="{int(size * 2)}"/></w:rPr>')
        xml = (f'<w:fldSimple {nsdecls("w")} w:instr=" {instr} "><w:r>{rpr}<w:t>{shown}</w:t></w:r></w:fldSimple>')
        par._p.append(parse_xml(xml))

    # ------------------------------------------------------------------ pages
    def render_page(self, pm):
        self.pm = pm
        nodes = pm.nodes
        if nodes:
            L = max(0.0, min(n.x0 for n in nodes))
            R = min(pm.W, max(n.x1 for n in nodes))
            if R - L < 100:
                L, R = 36.0, pm.W - 36.0
            top = max(0.0, min(n.top for n in nodes) + self.cfg.y_shift)
        else:
            L, R, top = 36.0, pm.W - 36.0, 36.0
        self.page_top = top
        self.page_bottom = max(0.0, pm.H - self.ftr_top + 1.0) if self.ftr_top else 0.0   # le pied de page réserve sa place
        self.secs.begin(start="nextPage", W=pm.W, H=pm.H, left=L, right=pm.W - R, top=top,
                        bottom=self.page_bottom, cols=None, hdr=self.hdr_dist, ftr=self.ftr_dist)
        ctr = Ctr(self.doc)
        if not nodes:
            self.spacer(ctr, 1)           # page blanche : on garde la page
        if pm.overlays:                    # tracés vectoriels : ancrés dans un paragraphe minuscule en tête de page
            hp = self.par(ctr)
            self._tiny_para(hp)
            for k, sh in enumerate(pm.overlays):
                self._anchor_shape(hp, sh, self._shape_id + k)
            self._shape_id += len(pm.overlays)
            top += 1.0
        self.render_nodes(ctr, nodes, L, R, top)
        self.pending_break = False

    def save(self, path):
        self.secs.finish()
        self.doc.save(path)
